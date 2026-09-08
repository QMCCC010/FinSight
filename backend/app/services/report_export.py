from __future__ import annotations

from io import BytesIO
import re


def _plain_markdown(value: str) -> str:
    value = re.sub(r"\[([^\]]+)]\(([^)]+)\)", r"\1（\2）", value)
    value = value.replace("**", "").replace("`", "")
    return value.strip()


def _rows(markdown: str) -> list[tuple[str, object]]:
    result: list[tuple[str, object]] = []
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1].strip()):
            table = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [_plain_markdown(cell) for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in cells):
                    table.append(cells)
                index += 1
            result.append(("table", table))
            continue
        if line.startswith("# "):
            result.append(("title", _plain_markdown(line[2:])))
        elif line.startswith("## "):
            result.append(("heading", _plain_markdown(line[3:])))
        elif line.startswith("> "):
            result.append(("quote", _plain_markdown(line[2:])))
        elif re.match(r"^\d+\.\s", line):
            result.append(("number", _plain_markdown(re.sub(r"^\d+\.\s*", "", line))))
        elif line.startswith(("- ", "* ")):
            result.append(("bullet", _plain_markdown(line[2:])))
        elif line.strip():
            result.append(("paragraph", _plain_markdown(line)))
        else:
            result.append(("space", ""))
        index += 1
    return result


def export_docx(title: str, markdown: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for kind, value in _rows(markdown):
        if kind == "title":
            paragraph = document.add_heading(str(value), level=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "heading":
            document.add_heading(str(value), level=1)
        elif kind == "quote":
            document.add_paragraph(str(value), style="Intense Quote")
        elif kind == "bullet":
            document.add_paragraph(str(value), style="List Bullet")
        elif kind == "number":
            document.add_paragraph(str(value), style="List Number")
        elif kind == "table":
            rows = value
            if not rows:
                continue
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Light Shading Accent 1"
            for row_index, row in enumerate(rows):
                for column_index, cell in enumerate(row):
                    table.cell(row_index, column_index).text = cell
        elif kind == "space":
            document.add_paragraph("")
        else:
            document.add_paragraph(str(value))
    core = document.core_properties
    core.title = title
    core.subject = "FinSight金融研究报告"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_pdf(title: str, markdown: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=title)
    styles = getSampleStyleSheet()
    base = ParagraphStyle("Chinese", parent=styles["BodyText"], fontName="STSong-Light", fontSize=9.5, leading=16, textColor=colors.HexColor("#26372f"), spaceAfter=5)
    heading = ParagraphStyle("ChineseHeading", parent=base, fontSize=15, leading=22, textColor=colors.HexColor("#173a2f"), spaceBefore=12, spaceAfter=8)
    title_style = ParagraphStyle("ChineseTitle", parent=heading, fontSize=22, leading=30, alignment=TA_CENTER, spaceAfter=18)
    quote = ParagraphStyle("ChineseQuote", parent=base, leftIndent=10, textColor=colors.HexColor("#66756d"), backColor=colors.HexColor("#f2f5f0"), borderPadding=6)
    story = []
    for kind, value in _rows(markdown):
        if kind == "table":
            rows = value
            if not rows:
                continue
            data = [[Paragraph(cell.replace("&", "&amp;").replace("<", "&lt;"), base) for cell in row] for row in rows]
            width = max(len(row) for row in rows)
            table = Table(data, colWidths=[(A4[0] - 36 * mm) / width] * width, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf0e5")),
                ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#cbd5cc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([table, Spacer(1, 7)])
            continue
        text = str(value).replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br/>")
        if kind == "title":
            story.append(Paragraph(text, title_style))
        elif kind == "heading":
            story.append(Paragraph(text, heading))
        elif kind == "quote":
            story.append(Paragraph(text, quote))
        elif kind == "bullet":
            story.append(Paragraph(f"• {text}", base))
        elif kind == "number":
            story.append(Paragraph(text, base))
        elif kind == "space":
            story.append(Spacer(1, 4))
        else:
            story.append(Paragraph(text, base))
    doc.build(story)
    return buffer.getvalue()
